# -*- coding: utf-8 -*-
"""Static layout audit: estimate each text frame's real ink height and flag
off-slide content, text-on-text collisions and text spilling out of its card."""
import os, sys
from pptx import Presentation
from pptx.util import Emu
from PIL import ImageFont

HERE = os.path.dirname(os.path.abspath(__file__))
PATH = os.path.join(HERE, "Xi'an - An Interesting Place.pptx")
FDIR = r"C:\Windows\Fonts"
SW, SH = 13.3333, 7.5

FONTS = {
    ("Georgia", False): "georgia.ttf", ("Georgia", True): "georgiab.ttf",
    ("Segoe UI", False): "segoeui.ttf", ("Segoe UI", True): "segoeuib.ttf",
    ("Segoe UI Semibold", False): "segoeuib.ttf", ("Segoe UI Semibold", True): "segoeuib.ttf",
    ("Microsoft YaHei", False): "msyh.ttc", ("Microsoft YaHei", True): "msyhbd.ttc",
}
_cache = {}


def _font(name, bold, size_pt):
    key = (name, bold, int(round(size_pt)))
    if key not in _cache:
        path = os.path.join(FDIR, FONTS.get((name, bold), "segoeui.ttf"))
        try:
            _cache[key] = ImageFont.truetype(path, max(8, int(round(size_pt * 4))))
        except Exception:
            _cache[key] = ImageFont.truetype(os.path.join(FDIR, "segoeui.ttf"),
                                             max(8, int(round(size_pt * 4))))
    return _cache[key]


def para_lines(p, width_in):
    """Wrapped line count, dominant font size and the widest line (inches)."""
    txt = "".join(r.text for r in p.runs)
    if not txt.strip():
        return 0, 10.0, 0.0
    size = max([(r.font.size.pt if r.font.size else 12) for r in p.runs] or [12])
    name = None
    bold = False
    for r in p.runs:
        if r.font.name:
            name = r.font.name
            bold = bool(r.font.bold)
            break
    f = _font(name or "Segoe UI", bold, size)
    max_w = width_in * 72.0 * 4.0
    lines, cur = [], ""
    for word in txt.split():
        trial = (cur + " " + word).strip()
        if f.getlength(trial) <= max_w or not cur:
            cur = trial
        else:
            lines.append(cur)
            cur = word
    lines.append(cur)
    widest = max(f.getlength(l) for l in lines) / 4.0 / 72.0
    return len(lines), size, widest


def ink_box(tf, width_in):
    """Estimated (height, width) of the ink actually laid down."""
    total, widest = 0.0, 0.0
    for p in tf.paragraphs:
        n, size, w = para_lines(p, width_in)
        if n == 0:
            continue
        sp = p.line_spacing if isinstance(p.line_spacing, float) else 1.2
        sb = p.space_before.pt if p.space_before else 0.0
        sa = p.space_after.pt if p.space_after else 0.0
        total += (sb + sa) / 72.0 + n * 1.2 * sp * size / 72.0
        widest = max(widest, w)
    return total, min(widest, width_in)


def rects(slide):
    texts, cards = [], []
    for sh in slide.shapes:
        try:
            l, t = sh.left / 914400.0, sh.top / 914400.0
            w, h = sh.width / 914400.0, sh.height / 914400.0
        except Exception:
            continue
        if sh.has_text_frame and sh.text_frame.text.strip():
            ih, iw = ink_box(sh.text_frame, w)
            texts.append((sh.text_frame.text.replace("\n", " | ")[:38], l, t, iw, ih))
        elif sh.shape_type is not None and w > 1.2 and h > 0.8:
            cards.append((l, t, w, h))
    return texts, cards


def overlap(a, b, pad=0.02):
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    return not (ax + aw <= bx + pad or bx + bw <= ax + pad or
                ay + ah <= by + pad or by + bh <= ay + pad)


def main():
    prs = Presentation(PATH)
    bad = 0
    for i, slide in enumerate(prs.slides, 1):
        texts, cards = rects(slide)
        msgs = []
        for txt, x, y, w, h in texts:
            if x < -0.01 or y < -0.01 or x + w > SW + 0.01 or y + h > SH + 0.01:
                msgs.append("OFF-SLIDE  y=%.2f h=%.2f  %r" % (y, h, txt))
            for (cl, ct, cw, ch) in cards:
                inside = (x >= cl - 0.05 and x <= cl + cw - 0.05 and
                          y >= ct - 0.05 and y <= ct + ch - 0.05)
                if inside and y + h > ct + ch + 0.02:
                    msgs.append("CARD-SPILL card(%.2f..%.2f) text y=%.2f h=%.2f -> %.2f  %r"
                                % (ct, ct + ch, y, h, y + h, txt))
        for a in range(len(texts)):
            for b in range(a + 1, len(texts)):
                ta, tb = texts[a], texts[b]
                if overlap((ta[1], ta[2], ta[3], ta[4]), (tb[1], tb[2], tb[3], tb[4])):
                    msgs.append("TEXT-OVERLAP %r  <->  %r" % (ta[0], tb[0]))
        if msgs:
            bad += len(msgs)
            print("\n--- slide %d ---" % i)
            for m in dict.fromkeys(msgs):
                print("   ", m)
    print("\ntotal issues:", bad)


if __name__ == "__main__":
    main()
