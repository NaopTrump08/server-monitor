# -*- coding: utf-8 -*-
"""Dump the paragraphs (and any tables) of a .docx so its content can be reviewed."""
import sys
from docx import Document

path = sys.argv[1]
doc = Document(path)

print("=== PARAGRAPHS ===")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t:
        print("[%02d][%s] %s" % (i, p.style.name if p.style else "-", t))

for ti, tbl in enumerate(doc.tables):
    print("\n=== TABLE %d (%d x %d) ===" % (ti, len(tbl.rows), len(tbl.columns)))
    for r in tbl.rows:
        print(" | ".join(c.text.strip().replace("\n", " ") for c in r.cells))
