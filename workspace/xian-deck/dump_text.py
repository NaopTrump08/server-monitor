# -*- coding: utf-8 -*-
"""Dump every piece of visible text plus the speaker notes, for vocabulary review."""
import os, re
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
prs = Presentation(os.path.join(HERE, "Xi'an - An Interesting Place.pptx"))

words = {}
for i, s in enumerate(prs.slides, 1):
    print("\n===== SLIDE %d =====" % i)
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            t = " | ".join(p.text for p in sh.text_frame.paragraphs if p.text.strip())
            print("  ", t)
    if s.has_notes_slide:
        nt = s.notes_slide.notes_text_frame.text
        print("   [NOTES]", nt)
        for w in re.findall(r"[A-Za-z'\-]+", nt):
            words[w.lower()] = words.get(w.lower(), 0) + 1

print("\n\n===== word list (notes, longest first) =====")
print(", ".join(sorted(words, key=len, reverse=True)))
