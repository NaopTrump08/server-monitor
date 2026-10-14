# -*- coding: utf-8 -*-
"""躺平与内卷——青年学生的为与不为（深色商务质感版）。"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------------- 深色商务配色 ----------------
BG_TOP    = RGBColor(0x1B, 0x22, 0x30)   # 渐变亮端（深蓝灰）
BG_BOT    = RGBColor(0x0D, 0x0F, 0x15)   # 渐变暗端（近黑）
PANEL     = RGBColor(0x1A, 0x1F, 0x2A)   # 卡片面板
PANEL_2   = RGBColor(0x20, 0x26, 0x33)   # 面板亮一点
EDGE      = RGBColor(0x34, 0x3C, 0x4C)   # 卡片描边
EDGE_SOFT = RGBColor(0x28, 0x2F, 0x3C)   # 更淡描边
GOLD      = RGBColor(0xC9, 0xA3, 0x44)   # 香槟金
GOLD_SOFT = RGBColor(0xE7, 0xD1, 0x92)   # 浅金
GOLD_DIM  = RGBColor(0x8A, 0x76, 0x48)   # 暗金
TEXT      = RGBColor(0xF0, 0xED, 0xE5)   # 主文字（米白）
MUTED     = RGBColor(0x9B, 0xA3, 0xB2)   # 次级文字
FAINT     = RGBColor(0x66, 0x6F, 0x80)   # 弱化文字
GREEN     = RGBColor(0x5F, 0xA8, 0x85)   # 「利」绿
GREEN_T   = RGBColor(0x16, 0x24, 0x20)   # 绿面板
GREEN_E   = RGBColor(0x2E, 0x4A, 0x3E)   # 绿描边
RED       = RGBColor(0xC6, 0x5D, 0x50)   # 「弊」红
RED_T     = RGBColor(0x26, 0x1A, 0x1C)   # 红面板
RED_E     = RGBColor(0x4C, 0x2E, 0x30)   # 红描边

FONT = "微软雅黑"
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
EMU_IN = 914400.0

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def set_run(run, text, size, bold=False, color=TEXT, font=FONT, spacing=None):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', font)
    if spacing is not None:
        rPr.set('spc', str(int(spacing * 100)))


def add_textbox(slide, x, y, w, h, anchor=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    if anchor is not None:
        tf.vertical_anchor = anchor
    return tb, tf


def solid_fill(shape, rgb, alpha_pct=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb
    if alpha_pct is not None and alpha_pct < 100:
        spPr = shape._element.spPr
        solidFill = spPr.find(qn('a:solidFill'))
        if solidFill is not None:
            srgbClr = solidFill.find(qn('a:srgbClr'))
            if srgbClr is not None:
                for a in srgbClr.findall(qn('a:alpha')):
                    srgbClr.remove(a)
                a = srgbClr.makeelement(qn('a:alpha'), {'val': str(int(alpha_pct * 1000))})
                srgbClr.append(a)


def gradient_fill(shape, c1, c2, angle=120.0):
    shape.fill.gradient()
    stops = shape.fill.gradient_stops
    stops[0].color.rgb = c1
    stops[0].position = 0.0
    stops[1].color.rgb = c2
    stops[1].position = 1.0
    try:
        shape.fill.gradient_angle = angle
    except Exception:
        pass


def shape(slide, kind, x, y, w, h, fill=None, alpha=None, line=None, line_w=None):
    sp = slide.shapes.add_shape(kind, x, y, w, h)
    if fill is not None:
        solid_fill(sp, fill, alpha)
    else:
        sp.fill.background()
    if line is not None:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w or 1)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def dark_card(slide, x, y, w, h, fill=PANEL, radius=0.06, line=EDGE, line_w=1.0):
    sp = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=line, line_w=line_w)
    try:
        sp.adjustments[0] = radius
    except Exception:
        pass
    return sp


def text_in_shape(sp, text, size, color, bold=False, spacing=None, anchor=MSO_ANCHOR.MIDDLE):
    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    set_run(p.add_run(), text, size, bold, color, spacing=spacing)
    return sp


def fill_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


# ---------------- 页头（深色） ----------------
def content_header(slide, title, kicker, page_no, total=12):
    # 右上角装饰大圆
    shape(slide, MSO_SHAPE.OVAL, Inches(10.9), Inches(-2.1), Inches(4.6), Inches(4.6), fill=GOLD, alpha=5)
    # 左下角装饰小圆
    shape(slide, MSO_SHAPE.OVAL, Inches(-1.0), Inches(6.4), Inches(2.6), Inches(2.6), fill=GOLD, alpha=4)

    # 眉题：金色小方块 + 金色字
    shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(0.55), Inches(0.14), Inches(0.14), fill=GOLD)
    _, ktf = add_textbox(slide, Inches(1.14), Inches(0.46), Inches(8), Inches(0.34))
    set_run(ktf.paragraphs[0].add_run(), kicker, 12, True, GOLD, spacing=2.4)

    # 主标题
    _, ttf = add_textbox(slide, Inches(0.88), Inches(0.92), Inches(11.6), Inches(0.8))
    set_run(ttf.paragraphs[0].add_run(), title, 28, True, TEXT)

    # 金色装饰线
    shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.78), Inches(1.9), Inches(0.03), fill=GOLD)
    shape(slide, MSO_SHAPE.DIAMOND, Inches(0.88), Inches(1.74), Inches(0.11), Inches(0.11), fill=GOLD_SOFT)

    # 底部细线（页脚）
    shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(7.14), Inches(11.53), Inches(0.012), fill=EDGE_SOFT)
    # 页码：03 金 + /12 灰
    _, ptf = add_textbox(slide, Inches(11.3), Inches(6.78), Inches(1.6), Inches(0.34))
    pp = ptf.paragraphs[0]
    pp.alignment = PP_ALIGN.RIGHT
    set_run(pp.add_run(), f"{page_no:02d}", 13, True, GOLD)
    set_run(pp.add_run(), f" / {total:02d}", 11, False, FAINT)


def banner_note(slide, y, runs, bg=PANEL, accent=GOLD, x=Inches(0.9), w=Inches(11.53), h=Inches(0.9)):
    dark_card(slide, x, y, w, h, fill=bg, radius=0.10, line=EDGE)
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, Inches(0.1), h, fill=accent)
    _, tf = add_textbox(slide, x + Inches(0.38), y + Inches(0.08), w - Inches(0.66), h - Inches(0.16),
                        anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    for txt, sz, bold, col in runs:
        set_run(p.add_run(), txt, sz, bold, col)


def bullet_row(slide, x, y, w, marker_kind, marker_color, text, size=14):
    if marker_kind == 'dot':
        shape(slide, MSO_SHAPE.OVAL, x, y + Inches(0.06), Inches(0.15), Inches(0.15), fill=marker_color)
    elif marker_kind == 'diamond':
        shape(slide, MSO_SHAPE.DIAMOND, x, y + Inches(0.05), Inches(0.16), Inches(0.16), fill=marker_color)
    else:  # square
        shape(slide, MSO_SHAPE.RECTANGLE, x, y + Inches(0.05), Inches(0.15), Inches(0.15), fill=marker_color)
    _, tf = add_textbox(slide, x + Inches(0.34), y, w - Inches(0.34), Inches(0.5))
    set_run(tf.paragraphs[0].add_run(), text, size, False, TEXT)


# ================= 第 1 页：封面 =================
s = prs.slides.add_slide(BLANK)
bg = shape(s, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H, fill=BG_BOT)
gradient_fill(bg, BG_TOP, BG_BOT, 118.0)
# 顶/底金色细线
shape(s, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.05), fill=GOLD)
shape(s, MSO_SHAPE.RECTANGLE, 0, Inches(7.45), SLIDE_W, Inches(0.05), fill=GOLD)
# 内框（hairline）
shape(s, MSO_SHAPE.RECTANGLE, Inches(0.28), Inches(0.28), Inches(12.77), Inches(6.94),
      fill=None, line=GOLD_DIM, line_w=1.0)
# 装饰半透明大圆
shape(s, MSO_SHAPE.OVAL, Inches(9.5), Inches(-2.7), Inches(6.0), Inches(6.0), fill=GOLD, alpha=7)
shape(s, MSO_SHAPE.OVAL, Inches(10.1), Inches(-2.0), Inches(4.8), Inches(4.8), fill=GOLD, alpha=5)
shape(s, MSO_SHAPE.OVAL, Inches(-2.4), Inches(5.5), Inches(5.0), Inches(5.0), fill=BG_TOP, alpha=60)
# 课程标签
_, ltf = add_textbox(s, Inches(1.35), Inches(1.35), Inches(9), Inches(0.5))
set_run(ltf.paragraphs[0].add_run(), "思想政治课 · 主题讨论", 14.5, True, GOLD, spacing=2.4)
# 主标题
_, ttf = add_textbox(s, Inches(1.28), Inches(2.15), Inches(11.2), Inches(1.7))
p = ttf.paragraphs[0]
set_run(p.add_run(), "躺平", 60, True, TEXT)
set_run(p.add_run(), "  与  ", 38, True, GOLD)
set_run(p.add_run(), "内卷", 60, True, TEXT)
# 副标题
_, stf = add_textbox(s, Inches(1.3), Inches(3.7), Inches(11.3), Inches(0.8))
set_run(stf.paragraphs[0].add_run(), "—— 青年学生的「为」与「不为」", 26, True, GOLD_SOFT, spacing=1.2)
# 金色装饰线
shape(s, MSO_SHAPE.RECTANGLE, Inches(1.32), Inches(4.85), Inches(2.4), Inches(0.04), fill=GOLD)
# 引导语
_, gtf = add_textbox(s, Inches(1.32), Inches(5.15), Inches(11.0), Inches(1.4))
gp = gtf.paragraphs[0]
set_run(gp.add_run(), "拒绝无效内卷  ·  不做消极躺平", 20, False, TEXT, spacing=1.2)
gp2 = gtf.add_paragraph()
gp2.space_before = Pt(12)
set_run(gp2.add_run(), "在可为之事上全力以赴，在不可为之事上保持清醒", 14, False, MUTED, spacing=0.8)


# ================= 第 2 页：目录 =================
s = prs.slides.add_slide(BLANK)
fill_bg(s, BG_BOT)
content_header(s, "议题导航", "CONTENTS", 2)

items = [
    ("01", "现象观察", "什么是「内卷」？什么是「躺平」？"),
    ("02", "追根溯源", "两种社会心态为何在青年中流行？"),
    ("03", "辩证审视", "内卷与躺平的利与弊、误区何在？"),
    ("04", "青年之「为」", "应当做什么——该担当、该奋斗之处"),
    ("05", "青年之「不为」", "不应做什么——该拒绝、该清醒之处"),
    ("06", "总结寄语", "在奋斗与从容之间，找到自己的人生节奏"),
]
shape(s, MSO_SHAPE.RECTANGLE, Inches(1.34), Inches(2.0), Inches(0.02), Inches(4.4), fill=EDGE)
y = Inches(1.92)
for i, (num, t, d) in enumerate(items):
    # 编号圆
    badge = shape(s, MSO_SHAPE.OVAL, Inches(1.02), y, Inches(0.64), Inches(0.64), fill=None, line=GOLD, line_w=1.2)
    text_in_shape(badge, num, 14, GOLD, True)
    # 标题 + 描述
    _, ttf = add_textbox(s, Inches(1.95), y - Inches(0.02), Inches(10.4), Inches(0.85))
    set_run(ttf.paragraphs[0].add_run(), t, 17.5, True, TEXT)
    dp = ttf.add_paragraph()
    dp.space_before = Pt(2)
    set_run(dp.add_run(), d, 12.5, False, MUTED)
    y += Inches(0.86)


# ================= 现象页（3、4） =================
def phenomenon_slide(page_no, title, kicker, concept, shows, judgement):
    s = prs.slides.add_slide(BLANK)
    fill_bg(s, BG_BOT)
    content_header(s, title, kicker, page_no)

    dark_card(s, Inches(0.9), Inches(2.1), Inches(11.53), Inches(1.12), radius=0.08)
    shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.1), Inches(0.1), Inches(1.12), fill=GOLD)
    _, ctf = add_textbox(s, Inches(1.28), Inches(2.22), Inches(10.9), Inches(0.92), anchor=MSO_ANCHOR.MIDDLE)
    cp = ctf.paragraphs[0]
    set_run(cp.add_run(), "概念解读   ", 14.5, True, GOLD)
    set_run(cp.add_run(), concept, 14.5, False, TEXT)

    _, htf = add_textbox(s, Inches(0.9), Inches(3.42), Inches(6), Inches(0.5))
    set_run(htf.paragraphs[0].add_run(), "常见表现", 16, True, GOLD_SOFT, spacing=1.4)
    y = Inches(3.95)
    for b in shows:
        bullet_row(s, Inches(1.05), y, Inches(11.0), 'diamond', GOLD, b)
        y += Inches(0.52)

    banner_note(s, Inches(6.05), [
        ("核心判断  ", 14, True, GOLD),
        (judgement, 14, False, TEXT),
    ])


phenomenon_slide(
    3, "现象观察 · 什么是「内卷」", "PHENOMENON 01",
    "「内卷」原指社会或组织发展到一定阶段后，无法向外拓展，只能在内部展开过度、低效的重复竞争的状态。",
    ["加班熬夜、无效内耗，比拼谁更“辛苦”",
     "同质化竞争，千军万马挤同一座独木桥",
     "“剧场效应”：前排站起，后排被迫站起"],
    "竞争本身并不可怕，可怕的是无意义、低效率、耗竭身心的“内耗式竞争”。",
)

phenomenon_slide(
    4, "现象观察 · 什么是「躺平」", "PHENOMENON 02",
    "「躺平」指部分年轻人面对压力，选择放弃过度竞争、主动降低欲望，追求低消耗、低期待生活的一种态度。",
    ["降低物质消费，追求“够用就好”",
     "拒绝无意义加班，主张工作与生活平衡",
     "佛系心态，对结果看得更淡、更从容"],
    "适度放松是自我调节，彻底放弃是逃避；关键在于“度”。",
)


# ================= 原因页（5、6） =================
def reason_slide(page_no, title, kicker, reasons, summary, accent=GOLD):
    s = prs.slides.add_slide(BLANK)
    fill_bg(s, BG_BOT)
    content_header(s, title, kicker, page_no)
    pos = [(0.9, 2.1), (6.85, 2.1), (0.9, 4.3), (6.85, 4.3)]
    for (x, y), (t, d), idx in zip(pos, reasons, [1, 2, 3, 4]):
        dark_card(s, Inches(x), Inches(y), Inches(5.58), Inches(2.0), radius=0.07)
        # 顶部金条
        shape(s, MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(5.58), Inches(0.08), fill=accent)
        # 序号
        shape(s, MSO_SHAPE.OVAL, Inches(x + 0.3), Inches(y + 0.32), Inches(0.5), Inches(0.5), fill=None, line=accent, line_w=1.2)
        _, ntf = add_textbox(s, Inches(x + 0.3), Inches(y + 0.34), Inches(0.5), Inches(0.46))
        np_ = ntf.paragraphs[0]
        np_.alignment = PP_ALIGN.CENTER
        set_run(np_.add_run(), f"0{idx}", 12.5, True, accent)
        # 标题 + 描述
        _, ttf = add_textbox(s, Inches(x + 0.95), Inches(y + 0.34), Inches(4.4), Inches(1.55))
        set_run(ttf.paragraphs[0].add_run(), t, 16.5, True, TEXT)
        dp = ttf.add_paragraph()
        dp.space_before = Pt(5)
        set_run(dp.add_run(), d, 12.5, False, MUTED)
    banner_note(s, Inches(6.35), [
        ("小结  ", 13, True, GOLD),
        (summary, 13, False, TEXT),
    ])


reason_slide(
    5, "追根溯源 · 内卷为何流行", "REASON 01", [
        ("资源有限", "优质岗位、升学机会相对稀缺，需求远超供给，形成竞争压力。"),
        ("标准单一", "评价体系趋于同质化，“分数”“排名”成为衡量价值的主要尺度。"),
        ("焦虑传导", "社会、家庭、同辈的焦虑相互放大，个体被迫卷入竞争漩涡。"),
        ("发展约束", "特定发展阶段内，增量空间收窄，竞争转向存量“内耗”。"),
    ],
    "内卷不是某一个人的问题，而是发展阶段、评价机制与群体心理共同作用的结果。",
    GOLD,
)

reason_slide(
    6, "追根溯源 · 躺平为何出现", "REASON 02", [
        ("对内卷的逆反", "当竞争变得无意义且过度时，部分人选择主动退出，以“躺”表达不满。"),
        ("高压下的自保", "长期高压透支身心，躺平成为一种心理上的自我保护与缓冲。"),
        ("反思单一成功", "对“唯成功论”的质疑，促使人们重新思考人生价值的多元可能。"),
        ("时代提供选择", "物质基础改善、价值取向多元，让“按自己的节奏生活”成为可能。"),
    ],
    "躺平既是社会心态的折射，也包含对人生意义与生活方式的一次认真追问。",
    GOLD,
)


# ================= 辩证页（7、8） =================
def dialectic_slide(page_no, title, kicker, goods, bads, conclusion):
    s = prs.slides.add_slide(BLANK)
    fill_bg(s, BG_BOT)
    content_header(s, title, kicker, page_no)

    dark_card(s, Inches(0.9), Inches(2.05), Inches(5.6), Inches(4.0), fill=GREEN_T, line=GREEN_E, radius=0.05)
    shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.05), Inches(5.6), Inches(0.08), fill=GREEN)
    _, htf = add_textbox(s, Inches(1.2), Inches(2.28), Inches(5.0), Inches(0.5))
    set_run(htf.paragraphs[0].add_run(), "积极一面（利）", 17, True, GREEN, spacing=1.2)
    y = Inches(2.95)
    for g in goods:
        bullet_row(s, Inches(1.2), y, Inches(4.95), 'dot', GREEN, g, 13.5)
        y += Inches(0.72)

    dark_card(s, Inches(6.85), Inches(2.05), Inches(5.6), Inches(4.0), fill=RED_T, line=RED_E, radius=0.05)
    shape(s, MSO_SHAPE.RECTANGLE, Inches(6.85), Inches(2.05), Inches(5.6), Inches(0.08), fill=RED)
    _, htf2 = add_textbox(s, Inches(7.15), Inches(2.28), Inches(5.0), Inches(0.5))
    set_run(htf2.paragraphs[0].add_run(), "消极一面（弊）", 17, True, RED, spacing=1.2)
    y = Inches(2.95)
    for b in bads:
        bullet_row(s, Inches(7.15), y, Inches(4.95), 'square', RED, b, 13.5)
        y += Inches(0.72)

    banner_note(s, Inches(6.3), [
        ("结论  ", 13, True, GOLD),
        (conclusion, 13, False, TEXT),
    ])


dialectic_slide(
    7, "辩证审视 · 内卷的利与弊", "DIALECTIC 01",
    ["适度竞争能够激发潜能、催人奋进",
     "推动个人成长与集体事业不断进步",
     "让努力与才华有更多被看见的机会"],
    ["过度内耗透支身心，侵蚀健康与生活",
     "挤占创造空间，陷入低水平重复",
     "催生焦虑攀比，扭曲奋斗的意义"],
    "我们要警惕的不是“努力”，而是“无效的内耗”。",
)

dialectic_slide(
    8, "辩证审视 · 躺平的利与弊", "DIALECTIC 02",
    ["提供喘息空间，及时调整身心状态",
     "反思人生意义，拒绝盲目跟风攀比",
     "倡导生活平衡，重视质量与幸福感"],
    ["彻底躺平丧失斗志，虚度大好青春",
     "以“佛系”之名逃避责任与成长",
     "放弃奋斗，最终失去选择的能力"],
    "躺平可以“歇脚”，但不能“就地长眠”；休息是为了更好地出发。",
)


# ================= 为 / 不为页（9、10） =================
def action_slide(page_no, title, kicker, items, accent, marker):
    s = prs.slides.add_slide(BLANK)
    fill_bg(s, BG_BOT)
    content_header(s, title, kicker, page_no)
    pos = [(0.9, 2.05), (6.85, 2.05), (0.9, 4.25), (6.85, 4.25)]
    for (x, y), (t, d), idx in zip(pos, items, [1, 2, 3, 4]):
        dark_card(s, Inches(x), Inches(y), Inches(5.58), Inches(2.05), radius=0.07)
        shape(s, MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.1), Inches(2.05), fill=accent)
        # 圆形图标
        shape(s, MSO_SHAPE.OVAL, Inches(x + 0.4), Inches(y + 0.32), Inches(0.56), Inches(0.56), fill=accent)
        _, ntf = add_textbox(s, Inches(x + 0.4), Inches(y + 0.33), Inches(0.56), Inches(0.54))
        np_ = ntf.paragraphs[0]
        np_.alignment = PP_ALIGN.CENTER
        set_run(np_.add_run(), f"0{idx}", 14, True, BG_BOT)
        # 标题 + 描述
        _, ttf = add_textbox(s, Inches(x + 1.14), Inches(y + 0.32), Inches(4.25), Inches(1.6))
        set_run(ttf.paragraphs[0].add_run(), t, 16.5, True, TEXT)
        dp = ttf.add_paragraph()
        dp.space_before = Pt(5)
        set_run(dp.add_run(), d, 12.5, False, MUTED)


action_slide(
    9, "青年之「为」· 应当做什么", "WHAT TO DO", [
        ("为理想而奋斗", "立大志、明大德，把个人小我融入国家大我，做有志向的青年。"),
        ("为成长而学习", "打牢知识根基、练就过硬本领，做有本领、能担当的青年。"),
        ("为担当而奉献", "在服务他人、奉献社会中实现价值，做有责任感的青年。"),
        ("为健康而自律", "强健体魄、磨炼意志、张弛有度，做有韧性的青年。"),
    ], GOLD, 'dot',
)

action_slide(
    10, "青年之「不为」· 不应做什么", "WHAT NOT TO DO", [
        ("不为无效内卷", "拒绝盲目跟风、虚耗时日的低效竞争，把力气用在刀刃上。"),
        ("不为消极躺平", "拒绝逃避责任、自暴自弃，不做“精致的利己”与“消极的看客”。"),
        ("不为攀比焦虑", "拒绝被单一标准绑架，不被分数、排名、面子牵着走。"),
        ("不为急功近利", "拒绝投机取巧、走捷径，脚踏实地走好每一步。"),
    ], RED, 'square',
)


# ================= 第 11 页：正确选择 =================
s = prs.slides.add_slide(BLANK)
fill_bg(s, BG_BOT)
content_header(s, "正确选择 · 在奋斗与从容间找到平衡", "THE RIGHT PATH", 11)

points = [
    "奋斗是青春最亮丽的底色——人生因奋斗而精彩。",
    "拒绝无效内卷，追求高效成长——让努力更有意义。",
    "避免消极躺平，保持向上姿态——让心态始终在线。",
    "找到自己的节奏，实现有意义的人生——不盲从、不虚度。",
]
y = Inches(2.0)
for i, ptext in enumerate(points):
    circle = shape(s, MSO_SHAPE.OVAL, Inches(1.0), y, Inches(0.55), Inches(0.55), fill=None, line=GOLD, line_w=1.4)
    text_in_shape(circle, str(i + 1), 15, GOLD, True)
    if i < len(points) - 1:
        shape(s, MSO_SHAPE.RECTANGLE, Inches(1.26), y + Inches(0.55), Inches(0.03), Inches(0.3), fill=EDGE)
    _, ttf = add_textbox(s, Inches(1.85), y - Inches(0.02), Inches(10.4), Inches(0.7))
    set_run(ttf.paragraphs[0].add_run(), ptext, 15, False, TEXT)
    y += Inches(0.84)

# 金句卡
dark_card(s, Inches(0.9), Inches(5.55), Inches(11.53), Inches(1.1), radius=0.08)
shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(5.55), Inches(0.1), Inches(1.1), fill=GOLD)
_, qtf = add_textbox(s, Inches(1.28), Inches(5.68), Inches(10.8), Inches(0.85), anchor=MSO_ANCHOR.MIDDLE)
set_run(qtf.paragraphs[0].add_run(), "青春由磨砺而出彩，人生因奋斗而升华。", 20, True, GOLD_SOFT, spacing=1.2)


# ================= 第 12 页：结语 =================
s = prs.slides.add_slide(BLANK)
bg = shape(s, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H, fill=BG_BOT)
gradient_fill(bg, BG_TOP, BG_BOT, 118.0)
shape(s, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.05), fill=GOLD)
shape(s, MSO_SHAPE.RECTANGLE, 0, Inches(7.45), SLIDE_W, Inches(0.05), fill=GOLD)
shape(s, MSO_SHAPE.RECTANGLE, Inches(0.28), Inches(0.28), Inches(12.77), Inches(6.94),
      fill=None, line=GOLD_DIM, line_w=1.0)
shape(s, MSO_SHAPE.OVAL, Inches(9.6), Inches(-2.7), Inches(6.0), Inches(6.0), fill=GOLD, alpha=7)
shape(s, MSO_SHAPE.OVAL, Inches(-2.2), Inches(5.5), Inches(4.8), Inches(4.8), fill=BG_TOP, alpha=60)

_, htf = add_textbox(s, Inches(1.3), Inches(1.25), Inches(11.0), Inches(0.6))
set_run(htf.paragraphs[0].add_run(), "结语 · 与青年共勉", 22, True, GOLD_SOFT, spacing=2.4)
shape(s, MSO_SHAPE.RECTANGLE, Inches(1.32), Inches(2.0), Inches(1.8), Inches(0.04), fill=GOLD)

_, btf = add_textbox(s, Inches(1.3), Inches(2.6), Inches(11.2), Inches(1.9))
lines = [
    "内卷不可怕，可怕的是迷失方向的盲目消耗；",
    "躺平不可取，可贵的是张弛有度的清醒选择。",
]
for i, ln in enumerate(lines):
    p = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
    p.space_after = Pt(14)
    set_run(p.add_run(), ln, 21, False, TEXT, spacing=0.6)

_, gtf = add_textbox(s, Inches(1.3), Inches(4.6), Inches(11.2), Inches(1.7))
gp = gtf.paragraphs[0]
gp.alignment = PP_ALIGN.CENTER
set_run(gp.add_run(), "愿每一位青年：", 18, True, GOLD_SOFT, spacing=2.4)
gp2 = gtf.add_paragraph()
gp2.alignment = PP_ALIGN.CENTER
gp2.space_before = Pt(12)
set_run(gp2.add_run(), "在可为之事上全力以赴，在不可为之事上保持清醒，", 16, False, TEXT, spacing=0.6)
gp3 = gtf.add_paragraph()
gp3.alignment = PP_ALIGN.CENTER
gp3.space_before = Pt(6)
set_run(gp3.add_run(), "拒绝无效内卷，不做消极躺平，书写无愧于时代的青春篇章。", 16, False, TEXT, spacing=0.6)

chip = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.17), Inches(6.3), Inches(7.0), Inches(0.62),
             fill=None, line=GOLD, line_w=1.2)
try:
    chip.adjustments[0] = 0.5
except Exception:
    pass
text_in_shape(chip, "课堂讨论：你如何看待身边的「内卷」与「躺平」？", 14, GOLD_SOFT, True, spacing=0.6)


out = "躺平与内卷——青年学生的为与不为.pptx"
prs.save(out)
print("saved:", out)
