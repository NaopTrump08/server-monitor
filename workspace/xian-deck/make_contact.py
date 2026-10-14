# -*- coding: utf-8 -*-
"""Contact sheet of every rendered page, plus a speaker-notes report."""
import os, glob
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "preview", "final")
OUT = os.path.join(HERE, "preview", "contact.png")

pages = sorted(glob.glob(os.path.join(SRC, "page-*.png")))
cols, tw = 3, 620
th = int(tw * 9 / 16)
pad, label = 18, 30
rows = (len(pages) + cols - 1) // cols
sheet = Image.new("RGB", (cols * (tw + pad) + pad, rows * (th + label + pad) + pad), (238, 234, 226))
d = ImageDraw.Draw(sheet)
f = ImageFont.truetype(r"C:\Windows\Fonts\segoeuib.ttf", 20)
for i, p in enumerate(pages):
    im = Image.open(p).convert("RGB").resize((tw, th), Image.LANCZOS)
    r, c = divmod(i, cols)
    x = pad + c * (tw + pad)
    y = pad + r * (th + label + pad)
    sheet.paste(im, (x, y))
    d.rectangle([x - 1, y - 1, x + tw, y + th], outline=(180, 172, 160))
    d.text((x + 2, y + th + 4), "page %d" % (i + 1), font=f, fill=(60, 55, 50))
sheet.save(OUT)
print("contact sheet ->", OUT, sheet.size)

prs = Presentation(os.path.join(HERE, "Xi'an - An Interesting Place.pptx"))
print("\nspeaker notes:")
for i, s in enumerate(prs.slides, 1):
    txt = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
    print("  slide %2d : %3d words | %s" % (i, len(txt.split()), txt[:64].replace("\n", " ")))
