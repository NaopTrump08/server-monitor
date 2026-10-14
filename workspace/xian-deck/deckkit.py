# -*- coding: utf-8 -*-
"""Layout kit for the Xi'an deck: warm-paper editorial system, no template defaults.

Type system:  Georgia (display / numerals)  +  Segoe UI (text / labels)
Palette:      drawn from the city itself - terracotta clay, city-wall ink,
              Tang gold, silk-road jade.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")

# ----------------------------------------------------------------- palette
INK        = (23, 25, 34)
INK_2      = (33, 36, 48)
SLATE      = (43, 47, 58)
PAPER      = (247, 242, 233)
PAPER_2    = (239, 230, 215)
PAPER_LINE = (214, 201, 180)
CLAY       = (192, 92, 60)
CLAY_DEEP  = (154, 69, 39)
CLAY_SOFT  = (223, 183, 159)
GOLD       = (199, 154, 69)
GOLD_SOFT  = (224, 192, 122)
JADE       = (62, 124, 105)
VERM       = (195, 58, 46)
MUTED      = (122, 114, 102)      # warm grey on paper
MUTED_D    = (150, 145, 156)      # grey on dark
DIM        = (168, 160, 148)

# ----------------------------------------------------------------- type
SERIF = "Georgia"
SANS  = "Segoe UI"
SANSB = "Segoe UI Semibold"
CJK   = "Microsoft YaHei"

# ----------------------------------------------------------------- geometry
SW, SH = 13.3333, 7.5
MARGIN = 0.82
CW = SW - 2 * MARGIN          # 11.6933
EYEBROW_Y = 0.60
TITLE_Y = 0.84
RULE_Y = 1.63
BODY_Y = 1.90
FOOT_Y = 7.06


def rgb(c):
    return RGBColor(*c)


def emu(v):
    return Emu(int(round(v * 914400)))


# ----------------------------------------------------------------- deck
def new_deck():
    prs = Presentation()
    prs.slide_width = emu(SW)
    prs.slide_height = emu(SH)
    return prs


def slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


# ----------------------------------------------------------------- shapes
def bg(sl, name, dark_hint=None):
    """Full-bleed baked background image."""
    return sl.shapes.add_picture(os.path.join(ASSETS, name), 0, 0, emu(SW), emu(SH))


def rect(sl, x, y, w, h, fill=None, line=None, lw=0.75, radius=None,
         dash=None, shape=MSO_SHAPE.RECTANGLE, rot=None):
    if radius is not None:
        shape = MSO_SHAPE.ROUNDED_RECTANGLE
    s = sl.shapes.add_shape(shape, emu(x), emu(y), emu(w), emu(h))
    s.shadow.inherit = False
    # drop the theme shape style: it re-introduces a default drop shadow in
    # LibreOffice/PowerPoint even when effectLst is empty.
    style = s._element.find(qn('p:style'))
    if style is not None:
        s._element.remove(style)
    if radius is not None:
        try:
            s.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(lw)
        if dash:
            s.line.dash_style = dash
    if rot:
        s.rotation = rot
    s.text_frame.word_wrap = True
    return s


def hairline(sl, x, y, w, color, lw=0.9, alpha=None):
    s = rect(sl, x, y, w, lw / 72.0, fill=color)
    return s


def vline(sl, x, y, h, color, lw=0.9):
    return rect(sl, x, y, lw / 72.0, h, fill=color)


def circle(sl, cx, cy, d, fill=None, line=None, lw=0.9):
    return rect(sl, cx - d / 2.0, cy - d / 2.0, d, d, fill=fill, line=line, lw=lw,
                shape=MSO_SHAPE.OVAL)


def pic(sl, name, x, y, w=None, h=None):
    """Place an asset preserving its aspect ratio (give w or h)."""
    from PIL import Image
    p = os.path.join(ASSETS, name)
    iw, ih = Image.open(p).size
    ar = iw / float(ih)
    if w is not None and h is None:
        h = w / ar
    elif h is not None and w is None:
        w = h * ar
    return sl.shapes.add_picture(p, emu(x), emu(y), emu(w), emu(h))


def pic_alpha(sl, name, x, y, w=None, h=None, alpha=0.30):
    """Place an asset with uniform transparency so it blends over any background."""
    p = pic(sl, name, x, y, w, h)
    p._element.blipFill.blip.append(parse_xml(
        '<a:alphaModFix xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
        ' amt="%d"/>' % int(round(alpha * 100000))))
    return p


def pic_right(sl, name, right, y, h, dark=False):
    """Place an asset so its right edge lands on `right`; returns its width."""
    g = pic(sl, name, 0, y, h=h)
    w = g.width / 914400.0
    g.left = emu(right - w)
    return w


# ----------------------------------------------------------------- text
def _style_run(r, size, font, color, bold, italic, spc, cjk_font=None):
    f = r.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    f.name = font
    rPr = r._r.get_or_add_rPr()
    if cjk_font:
        for tag in ("a:ea", "a:cs"):
            el = rPr.find(qn(tag))
            if el is None:
                el = parse_xml('<a:%s xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'
                               % tag.split(":")[1])
                rPr.append(el)
            el.set("typeface", cjk_font)
    if spc:
        rPr.set("spc", str(int(round(spc * 100))))


def textbox(sl, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = sl.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, runs, size=12.5, font=SANS, color=INK, bold=False, italic=False,
         spc=0, align=PP_ALIGN.LEFT, before=0, after=0, line=1.25,
         first=False, cjk_font=None, indent=None):
    """runs: str, or list of str / (text, {overrides})."""
    p = tf.paragraphs[0] if (first or (len(tf.paragraphs) == 1 and
                                       not tf.paragraphs[0].runs and
                                       tf.paragraphs[0].text == "")) else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = line
    if before:
        p.space_before = Pt(before)
    if after:
        p.space_after = Pt(after)
    if indent is not None:
        p.paragraph_format  # noqa
    if isinstance(runs, str):
        runs = [runs]
    for item in runs:
        if isinstance(item, str):
            t, ov = item, {}
        else:
            t, ov = item
        r = p.add_run()
        r.text = t
        _style_run(r, ov.get("size", size), ov.get("font", font),
                   ov.get("color", color), ov.get("bold", bold),
                   ov.get("italic", italic), ov.get("spc", spc),
                   ov.get("cjk", cjk_font))
    return p


def block(sl, x, y, w, h, blocks, anchor=MSO_ANCHOR.TOP):
    """blocks: list of dicts accepted by para()."""
    tf = textbox(sl, x, y, w, h, anchor)
    for i, b in enumerate(blocks):
        b = dict(b)
        b["first"] = (i == 0)
        para(tf, **b)
    return tf


# ----------------------------------------------------------------- chrome
def eyebrow(sl, label, color=CLAY, dash_color=None, x=MARGIN, y=EYEBROW_Y, w=None):
    """letterspaced small-caps label with a short leading rule."""
    rect(sl, x, y + 0.075, 0.30, 0.035, fill=dash_color or color)
    tf = textbox(sl, x + 0.44, y - 0.035, (w or CW) - 0.44, 0.34)
    para(tf, label.upper(), size=10.5, font=SANSB, color=color, spc=2.6, line=1.0, first=True)
    return tf


def header(sl, eyebrow_label, title, dark=False, sub=None, accent=CLAY,
           title_size=33, rule=True, title_w=None):
    fg = PAPER if dark else INK
    eyebrow(sl, eyebrow_label, color=GOLD if dark else accent)
    tw = title_w or (CW * 0.90)
    tf = textbox(sl, MARGIN, TITLE_Y, tw, 0.70)
    para(tf, title, size=title_size, font=SERIF, bold=True, color=fg, line=1.03, first=True)
    if sub:
        stf = textbox(sl, MARGIN, 1.44, min(tw + 0.8, CW), 0.54)
        para(stf, sub, size=13, font=SANS, color=(MUTED_D if dark else MUTED),
             line=1.34, first=True)
    if rule:
        hairline(sl, MARGIN, 2.02 if sub else 1.62, CW,
                 (58, 62, 76) if dark else PAPER_LINE, lw=1.1)


def tag(sl, x, y, text, color=CLAY, w=3.4):
    """small square bullet + letterspaced tag, used inside cards."""
    rect(sl, x, y + 0.055, 0.085, 0.085, fill=color)
    tf = textbox(sl, x + 0.21, y - 0.012, w, 0.28)
    para(tf, text.upper(), size=9, font=SANSB, color=color, spc=2.2, line=1.0, first=True)


def footer(sl, num, dark=False, label="XI'AN  \u00b7  AN INTERESTING PLACE"):
    fg = MUTED_D if dark else DIM
    tf = textbox(sl, MARGIN, FOOT_Y, CW * 0.7, 0.26)
    para(tf, label, size=8, font=SANS, color=fg, spc=2.2, line=1.0, first=True)
    ntf = textbox(sl, SW - MARGIN - 1.2, FOOT_Y - 0.015, 1.2, 0.28)
    para(ntf, "%02d" % num, size=10, font=SERIF, bold=True, color=(GOLD if dark else CLAY),
         align=PP_ALIGN.RIGHT, line=1.0, first=True)


def seal(sl, x, y, size=0.62, chars="\u897f\u5b89", fill=VERM, fg=PAPER, font_size=17):
    """vermilion seal stamp with knocked-out characters."""
    rect(sl, x, y, size, size, fill=fill, radius=0.14)
    tf = textbox(sl, x, y + 0.02, size, size - 0.04, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, chars, size=font_size, font=CJK, bold=True, color=fg,
         align=PP_ALIGN.CENTER, line=0.95, first=True, cjk_font=CJK)


def notes(sl, text):
    sl.notes_slide.notes_text_frame.text = text


# ----------------------------------------------------------------- reusable pieces
def card(sl, x, y, w, h, fill=PAPER_2, line=PAPER_LINE, radius=0.055, lw=0.9):
    return rect(sl, x, y, w, h, fill=fill, line=line, lw=lw, radius=radius)


def chip(sl, x, y, text, fill=None, fg=CLAY, line=None, w=None, size=9,
         font=SANSB, spc=2.0, padx=0.16, h=0.30, align=PP_ALIGN.CENTER):
    tw = w if w is not None else (len(text) * size * 0.0092 + spc * len(text) / 72.0 + padx * 2)
    s = rect(sl, x, y, tw, h, fill=fill, line=line, lw=0.9,
             radius=0.5 if fill else None, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = textbox(sl, x, y + 0.005, tw, h - 0.01, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, text, size=size, font=font, color=fg, spc=spc, align=align, line=1.0, first=True)
    return tw
