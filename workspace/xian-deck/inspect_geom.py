# -*- coding: utf-8 -*-
"""Dump real geometry of a few slides to verify the layout matches the design intent."""
import os, sys
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
PATH = os.path.join(HERE, "Xi'an - An Interesting Place.pptx")

prs = Presentation(PATH)
print("slide size (in): %.4f x %.4f" % (prs.slide_width / 914400.0, prs.slide_height / 914400.0))

for idx in [1, 2]:
    print("\n=== slide %d ===" % (idx + 1))
    for sh in list(prs.slides[idx].shapes)[:9]:
        t = ""
        if sh.has_text_frame:
            t = sh.text_frame.text.replace("\n", " / ")[:44]
        print("  x=%7.3f y=%7.3f w=%7.3f h=%7.3f  %s" % (
            sh.left / 914400.0, sh.top / 914400.0,
            sh.width / 914400.0, sh.height / 914400.0, t))
