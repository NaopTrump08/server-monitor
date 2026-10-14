# -*- coding: utf-8 -*-
"""生成思想政治课讨论"躺平"与"内卷"主题 PPT。"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------- 主题色 ----------
RED      = RGBColor(0xB0, 0x2E, 0x2E)   # 深红（主色）
RED_DARK = RGBColor(0x8C, 0x1F, 0x1F)   # 更深的红
GOLD     = RGBColor(0xD8, 0xA2, 0x3B)   # 金色点缀
GOLD_LT  = RGBColor(0xF0, 0xD9, 0xA0)
INK      = RGBColor(0x2B, 0x2B, 0x2B)   # 正文深灰
GREY     = RGBColor(0x6B, 0x6B, 0x6B)
CREAM    = RGBColor(0xFB, 0xF6, 0xEE)   # 米白背景
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT    = RGBColor(0xF4, 0xEC, 0xE4)   # 浅色卡片

FONT = "微软雅黑"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def set_run(run, text, size, bold=False, color=INK, font=FONT):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', font)


def add_textbox(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    return tb, tf


def fill_bg(slide, color):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def add_rect(slide, x, y, w, h, color, line=False, line_color=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line:
        shp.line.color.rgb = line_color
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_card(slide, x, y, w, h, fill=LIGHT):
    return add_rect(slide, x, y, w, h, fill)


# ---------- 通用页头 ----------
def content_header(slide, title, kicker, page_no):
    """内容页顶部标题区。"""
    # 顶部金色细条
    add_rect(slide, Inches(0.7), Inches(0.42), Inches(0.55), Inches(0.06), GOLD)
    # 眉题
    _, ktf = add_textbox(slide, Inches(0.7), Inches(0.52), Inches(9), Inches(0.3))
    kp = ktf.paragraphs[0]
    set_run(kp.add_run(), kicker, 12, True, GOLD)
    # 主标题
    _, ttf = add_textbox(slide, Inches(0.68), Inches(0.82), Inches(11.5), Inches(0.75))
    tp = ttf.paragraphs[0]
    set_run(tp.add_run(), title, 28, True, RED_DARK)
    # 页码
    _, ptf = add_textbox(slide, Inches(12.35), Inches(6.9), Inches(0.7), Inches(0.4))
    pp = ptf.paragraphs[0]
    pp.alignment = PP_ALIGN.RIGHT
    set_run(pp.add_run(), str(page_no), 12, False, GREY)
    # 底部金色细条
    add_rect(slide, Inches(0.7), Inches(7.06), Inches(0.55), Inches(0.04), GOLD)


# ---------- 第 1 页：封面 ----------
s = prs.slides.add_slide(BLANK)
fill_bg(s, RED)
# 顶部金色装饰线
add_rect(s, Inches(0), Inches(0), SLIDE_W, Inches(0.16), GOLD)
add_rect(s, Inches(0), Inches(7.34), SLIDE_W, Inches(0.16), GOLD)
# 左侧竖向装饰块
add_rect(s, Inches(0), Inches(0.16), Inches(0.28), Inches(7.18), RED_DARK)
# 课程标签
_, ltf = add_textbox(s, Inches(1.0), Inches(1.05), Inches(6), Inches(0.5))
lp = ltf.paragraphs[0]
set_run(lp.add_run(), "思想政治课 · 主题讨论", 16, True, GOLD_LT)
# 主标题
_, ttf = add_textbox(s, Inches(0.95), Inches(1.85), Inches(11.4), Inches(1.6))
tp = ttf.paragraphs[0]
set_run(tp.add_run(), "躺平 与 内卷", 58, True, WHITE)
# 副标题
_, stf = add_textbox(s, Inches(1.0), Inches(3.35), Inches(11.4), Inches(0.9))
sp = stf.paragraphs[0]
set_run(sp.add_run(), "—— 青年学生的「为」与「不为」", 28, True, GOLD_LT)
# 装饰分隔线
add_rect(s, Inches(1.0), Inches(4.55), Inches(2.6), Inches(0.05), GOLD)
# 引导语
_, gtf = add_textbox(s, Inches(1.0), Inches(4.85), Inches(10.5), Inches(1.2))
gp = gtf.paragraphs[0]
set_run(gp.add_run(), "拒绝无效内卷 · 不做消极躺平", 20, False, WHITE)
gp2 = gtf.add_paragraph()
gp2.space_before = Pt(10)
set_run(gp2.add_run(), "在可为之事上全力以赴，在不可为之事上保持清醒", 15, False, GOLD_LT)


# ---------- 第 2 页：目录 ----------
s = prs.slides.add_slide(BLANK)
fill_bg(s, CREAM)
content_header(s, "议题导航", "CONTENTS", 2)

items = [
    ("01", "现象观察", "什么是「内卷」？什么是「躺平」？"),
    ("02", "追根溯源", "两种社会心态为何在青年中流行？"),
    ("03", "辩证审视", "内卷与躺平的利与弊、误区何在？"),
    ("04", "青年之「为」", "应当做什么——该担当、该奋斗之处"),
    ("05", "青年之「不为」", "不应做什么——该拒绝、该清醒之处"),
    ("06", "总结寄语", "在奋斗与从容之间，找到自己的人生节奏"),
]
y = Inches(1.75)
for num, t, d in items:
    x0 = Inches(0.9)
    # 编号圆
    badge = s.shapes.add_shape(MSO_SHAPE.OVAL, x0, y, Inches(0.6), Inches(0.6))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RED
    badge.line.fill.background()
    badge.shadow.inherit = False
    btf = badge.text_frame
    btf.word_wrap = False
    bp = btf.paragraphs[0]
    bp.alignment = PP_ALIGN.CENTER
    set_run(bp.add_run(), num, 14, True, WHITE)
    # 标题 + 描述
    _, ttf = add_textbox(s, Inches(1.75), y - Inches(0.06), Inches(10.6), Inches(0.85))
    tp = ttf.paragraphs[0]
    set_run(tp.add_run(), t, 18, True, INK)
    dp = ttf.add_paragraph()
    dp.space_before = Pt(2)
    set_run(dp.add_run(), d, 13, False, GREY)
    y += Inches(0.88)


# ---------- 第 3 页：什么是内卷 ----------
s = prs.slides.add_slide(BLANK)
fill_bg(s, CREAM)
content_header(s, "现象观察 · 什么是「内卷」", "PHENOMENON 1", 3)

# 概念卡片
add_card(s, Inches(0.9), Inches(1.7), Inches(11.5), Inches(1.15), WHITE)
add_rect(s, Inches(0.9), Inches(1.7), Inches(0.12), Inches(1.15), GOLD)
_, ctf = add_textbox(s, Inches(1.25), Inches(1.85), Inches(10.9), Inches(0.9))
cp = ctf.paragraphs[0]
set_run(cp.add_run(), "概念解读：", 15, True, RED)
set_run(cp.add_run(), "「内卷」原指社会或组织发展到一定阶段后，无法向外拓展，只能在内部展开过度、低效的重复竞争的状态。", 15, False, INK)

# 表现
_, htf = add_textbox(s, Inches(0.9), Inches(3.1), Inches(6), Inches(0.5))
set_run(htf.paragraphs[0].add_run(), "常见表现", 16, True, RED_DARK)
bullets = [
    "加班熬夜、无效内耗，比拼谁更“辛苦”",
    "同质化竞争，千军万马挤同一座独木桥",
    "“剧场效应”：前排站起，后排被迫站起",
]
_, btf = add_textbox(s, Inches(1.1), Inches(3.55), Inches(11.0), Inches(1.4))
for i, b in enumerate(bullets):
    p = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
    p.space_after = Pt(6)
    set_run(p.add_run(), "▪  ", 14, True, GOLD)
    set_run(p.add_run(), b, 15, False, INK)

# 关键点
add_card(s, Inches(0.9), Inches(5.55), Inches(11.5), Inches(1.0), RGBColor(0xF7, 0xE6, 0xE0))
_, ktf = add_textbox(s, Inches(1.15), Inches(5.72), Inches(11.0), Inches(0.7))
kp = ktf.paragraphs[0]
set_run(kp.add_run(), "核心判断：", 15, True, RED)
set_run(kp.add_run(), "竞争本身并不可怕，可怕的是无意义、低效率、耗竭身心的“内耗式竞争”。", 15, False, INK)


# ---------- 第 4 页：什么是躺平 ----------
s = prs.slides.add_slide(BLANK)
fill_bg(s, CREAM)
content_header(s, "现象观察 · 什么是「躺平」", "PHENOMENON 2", 4)

add_card(s, Inches(0.9), Inches(1.7), Inches(11.5), Inches(1.15), WHITE)
add_rect(s, Inches(0.9), Inches(1.7), Inches(0.12), Inches(1.15), GOLD)
_, ctf = add_textbox(s, Inches(1.25), Inches(1.85), Inches(10.9), Inches(0.9))
cp = ctf.paragraphs[0]
set_run(cp.add_run(), "概念解读：", 15, True, RED)
set_run(cp.add_run(), "「躺平」指部分年轻人面对压力，选择放弃过度竞争、主动降低欲望，追求低消耗、低期待生活的一种态度。", 15, False, INK)

_, htf = add_textbox(s, Inches(0.9), Inches(3.1), Inches(6), Inches(0.5))
set_run(htf.paragraphs[0].add_run(), "常见表现", 16, True, RED_DARK)
bullets = [
    "降低物质消费，追求“够用就好”",
    "拒绝无意义加班，主张工作与生活平衡",
    "佛系心态，对结果看得更淡、更从容",
]
_, btf = add_textbox(s, Inches(1.1), Inches(3.55), Inches(11.0), Inches(1.4))
for i, b in enumerate(bullets):
    p = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
    p.space_after = Pt(6)
    set_run(p.add_run(), "▪  ", 14, True, GOLD)
    set_run(p.add_run(), b, 15, False, INK)

add_card(s, Inches(0.9), Inches(5.55), Inches(11.5), Inches(1.0), RGBColor(0xF7, 0xE6, 0xE0))
_, ktf = add_textbox(s, Inches(1.15), Inches(5.72), Inches(11.0), Inches(0.7))
kp = ktf.paragraphs[0]
set_run(kp.add_run(), "核心判断：", 15, True, RED)
set_run(kp.add_run(), "适度放松是自我调节，彻底放弃是逃避；关键在于“度”。", 15, False, INK)


# ---------- 第 5 页：内卷为何流行 ----------
s = prs.slides.add_slide(BLANK)
fill_bg(s, CREAM)
content_header(s, "追根溯源 · 内卷为何流行", "REASON 1", 5)

reasons = [
    ("资源有限", "优质岗位、升学机会相对稀缺，需求远超供给，形成竞争压力。"),
    ("标准单一", "评价体系趋于同质化，“分数”“排名”成为衡量价值的主要尺度。"),
    ("焦虑传导", "社会、家庭、同辈的焦虑相互放大，个体被迫卷入竞争漩涡。"),
    ("发展约束", "特定发展阶段内，增量空间收窄，竞争转向存量“内耗”。"),
]
pos = [(0.9, 1.75), (6.8, 1.75), (0.9, 3.9), (6.8, 3.9)]
for (x, y), (t, d) in zip(pos, reasons):
    add_card(s, Inches(x), Inches(y), Inches(5.6), Inches(1.95), WHITE)
    add_rect(s, Inches(x), Inches(y), Inches(5.6), Inches(0.14), RED)
    _, ttf = add_textbox(s, Inches(x + 0.3), Inches(y + 0.35), Inches(5.0), Inches(0.5))
    set_run(ttf.paragraphs[0].add_run(), t, 18, True, RED_DARK)
    _, dtf = add_textbox(s, Inches(x + 0.3), Inches(y + 0.95), Inches(5.0), Inches(0.9))
    set_run(dtf.paragraphs[0].add_run(), d, 13.5, False, INK)

_, btf = add_textbox(s, Inches(0.9), Inches(6.1), Inches(11.5), Inches(0.7))
bp = btf.paragraphs[0]
set_run(bp.add_run(), "小结：", 14, True, GOLD)
set_run(bp.add_run(), "内卷不是某一个人的问题，而是发展阶段、评价机制与群体心理共同作用的结果。", 14, False, INK)


# ---------- 第 6 页：躺平为何出现 ----------
s = prs.slides.add_slide(BLANK)
fill_bg(s, CREAM)
content_header(s, "追根溯源 · 躺平为何出现", "REASON 2", 6)

reasons = [
    ("对内卷的逆反", "当竞争变得无意义且过度时，部分人选择主动退出，以“躺”表达不满。"),
    ("高压下的自保", "长期高压透支身心，躺平成为一种心理上的自我保护与缓冲。"),
    ("反思单一成功", "对“唯成功论”的质疑，促使人们重新思考人生价值的多元可能。"),
    ("时代提供选择", "物质基础改善、价值取向多元，让“按自己的节奏生活”成为可能。"),
]
pos = [(0.9, 1.75), (6.8, 1.75), (0.9, 3.9), (6.8, 3.9)]
for (x, y), (t, d) in zip(pos, reasons):
    add_card(s, Inches(x), Inches(y), Inches(5.6), Inches(1.95), WHITE)
    add_rect(s, Inches(x), Inches(y), Inches(5.6), Inches(0.14), GOLD)
    _, ttf = add_textbox(s, Inches(x + 0.3), Inches(y + 0.35), Inches(5.0), Inches(0.5))
    set_run(ttf.paragraphs[0].add_run(), t, 18, True, RED_DARK)
    _, dtf = add_textbox(s, Inches(x + 0.3), Inches(y + 0.95), Inches(5.0), Inches(0.9))
    set_run(dtf.paragraphs[0].add_run(), d, 13.5, False, INK)

_, btf = add_textbox(s, Inches(0.9), Inches(6.1), Inches(11.5), Inches(0.7))
bp = btf.paragraphs[0]
set_run(bp.add_run(), "小结：", 14, True, GOLD)
set_run(bp.add_run(), "躺平既是社会心态的折射，也包含对人生意义与生活方式的一次认真追问。", 14, False, INK)


# ---------- 第 7 页：辩证审视 内卷 ----------
s = prs.slides.add_slide(BLANK)
fill_bg(s, CREAM)
content_header(s, "辩证审视 · 内卷的利与弊", "DIALECTIC 1", 7)

# 利
add_card(s, Inches(0.9), Inches(1.75), Inches(5.6), Inches(4.1), WHITE)
add_rect(s, Inches(0.9), Inches(1.75), Inches(5.6), Inches(0.14), RGBColor(0x2E, 0x7D, 0x5B))
_, htf = add_textbox(s, Inches(1.2), Inches(2.0), Inches(5.0), Inches(0.5))
set_run(htf.paragraphs[0].add_run(), "积极一面（利）", 18, True, RGBColor(0x2E, 0x7D, 0x5B))
goods = [
    "适度竞争能够激发潜能、催人奋进",
    "推动个人成长与集体事业不断进步",
    "让努力与才华有更多被看见的机会",
]
_, gtf = add_textbox(s, Inches(1.2), Inches(2.6), Inches(5.0), Inches(3.0))
for i, g in enumerate(goods):
    p = gtf.paragraphs[0] if i == 0 else gtf.add_paragraph()
    p.space_after = Pt(10)
    set_run(p.add_run(), "✓  ", 15, True, RGBColor(0x2E, 0x7D, 0x5B))
    set_run(p.add_run(), g, 14.5, False, INK)

# 弊
add_card(s, Inches(6.8), Inches(1.75), Inches(5.6), Inches(4.1), WHITE)
add_rect(s, Inches(6.8), Inches(1.75), Inches(5.6), Inches(0.14), RED)
_, htf2 = add_textbox(s, Inches(7.1), Inches(2.0), Inches(5.0), Inches(0.5))
set_run(htf2.paragraphs[0].add_run(), "消极一面（弊）", 18, True, RED)
bads = [
    "过度内耗透支身心，侵蚀健康与生活",
    "挤占创造空间，陷入低水平重复",
    "催生焦虑攀比，扭曲奋斗的意义",
]
_, btf2 = add_textbox(s, Inches(7.1), Inches(2.6), Inches(5.0), Inches(3.0))
for i, b in enumerate(bads):
    p = btf2.paragraphs[0] if i == 0 else btf2.add_paragraph()
    p.space_after = Pt(10)
    set_run(p.add_run(), "✗  ", 15, True, RED)
    set_run(p.add_run(), b, 14.5, False, INK)

_, btf = add_textbox(s, Inches(0.9), Inches(6.05), Inches(11.5), Inches(0.8))
bp = btf.paragraphs[0]
set_run(bp.add_run(), "结论：", 14, True, GOLD)
set_run(bp.add_run(), "我们要警惕的不是“努力”，而是“无效的内耗”。", 14, False, INK)


# ---------- 第 8 页：辩证审视 躺平 ----------
s = prs.slides.add_slide(BLANK)
fill_bg(s, CREAM)
content_header(s, "辩证审视 · 躺平的利与弊", "DIALECTIC 2", 8)

add_card(s, Inches(0.9), Inches(1.75), Inches(5.6), Inches(4.1), WHITE)
add_rect(s, Inches(0.9), Inches(1.75), Inches(5.6), Inches(0.14), RGBColor(0x2E, 0x7D, 0x5B))
_, htf = add_textbox(s, Inches(1.2), Inches(2.0), Inches(5.0), Inches(0.5))
set_run(htf.paragraphs[0].add_run(), "积极一面（利）", 18, True, RGBColor(0x2E, 0x7D, 0x5B))
goods = [
    "提供喘息空间，及时调整身心状态",
    "反思人生意义，拒绝盲目跟风攀比",
    "倡导生活平衡，重视质量与幸福感",
]
_, gtf = add_textbox(s, Inches(1.2), Inches(2.6), Inches(5.0), Inches(3.0))
for i, g in enumerate(goods):
    p = gtf.paragraphs[0] if i == 0 else gtf.add_paragraph()
    p.space_after = Pt(10)
    set_run(p.add_run(), "✓  ", 15, True, RGBColor(0x2E, 0x7D, 0x5B))
    set_run(p.add_run(), g, 14.5, False, INK)

add_card(s, Inches(6.8), Inches(1.75), Inches(5.6), Inches(4.1), WHITE)
add_rect(s, Inches(6.8), Inches(1.75), Inches(5.6), Inches(0.14), RED)
_, htf2 = add_textbox(s, Inches(7.1), Inches(2.0), Inches(5.0), Inches(0.5))
set_run(htf2.paragraphs[0].add_run(), "消极一面（弊）", 18, True, RED)
bads = [
    "彻底躺平丧失斗志，虚度大好青春",
    "以“佛系”之名逃避责任与成长",
    "放弃奋斗，最终失去选择的能力",
]
_, btf2 = add_textbox(s, Inches(7.1), Inches(2.6), Inches(5.0), Inches(3.0))
for i, b in enumerate(bads):
    p = btf2.paragraphs[0] if i == 0 else btf2.add_paragraph()
    p.space_after = Pt(10)
    set_run(p.add_run(), "✗  ", 15, True, RED)
    set_run(p.add_run(), b, 14.5, False, INK)

_, btf = add_textbox(s, Inches(0.9), Inches(6.05), Inches(11.5), Inches(0.8))
bp = btf.paragraphs[0]
set_run(bp.add_run(), "结论：", 14, True, GOLD)
set_run(bp.add_run(), "躺平可以“歇脚”，但不能“就地长眠”；休息是为了更好地出发。", 14, False, INK)


# ---------- 第 9 页：青年之「为」 ----------
s = prs.slides.add_slide(BLANK)
fill_bg(s, CREAM)
content_header(s, "青年之「为」· 应当做什么", "WHAT TO DO", 9)

do_items = [
    ("为理想而奋斗", "立大志、明大德，把个人小我融入国家大我，做有志向的青年。"),
    ("为成长而学习", "打牢知识根基、练就过硬本领，做有本领、能担当的青年。"),
    ("为担当而奉献", "在服务他人、奉献社会中实现价值，做有责任感的青年。"),
    ("为健康而自律", "强健体魄、磨炼意志、张弛有度，做有韧性的青年。"),
]
pos = [(0.9, 1.75), (6.8, 1.75), (0.9, 4.0), (6.8, 4.0)]
for (x, y), (t, d) in zip(pos, do_items):
    add_card(s, Inches(x), Inches(y), Inches(5.6), Inches(2.0), WHITE)
    add_rect(s, Inches(x), Inches(y), Inches(0.14), Inches(2.0), RED)
    _, ttf = add_textbox(s, Inches(x + 0.45), Inches(y + 0.3), Inches(4.9), Inches(0.5))
    set_run(ttf.paragraphs[0].add_run(), t, 18, True, RED_DARK)
    _, dtf = add_textbox(s, Inches(x + 0.45), Inches(y + 0.85), Inches(4.95), Inches(1.0))
    set_run(dtf.paragraphs[0].add_run(), d, 13.5, False, INK)


# ---------- 第 10 页：青年之「不为」 ----------
s = prs.slides.add_slide(BLANK)
fill_bg(s, CREAM)
content_header(s, "青年之「不为」· 不应做什么", "WHAT NOT TO DO", 10)

not_items = [
    ("不为无效内卷", "拒绝盲目跟风、虚耗时日的低效竞争，把力气用在刀刃上。"),
    ("不为消极躺平", "拒绝逃避责任、自暴自弃，不做“精致的利己”与“消极的看客”。"),
    ("不为攀比焦虑", "拒绝被单一标准绑架，不被分数、排名、面子牵着走。"),
    ("不为急功近利", "拒绝投机取巧、走捷径，脚踏实地走好每一步。"),
]
pos = [(0.9, 1.75), (6.8, 1.75), (0.9, 4.0), (6.8, 4.0)]
for (x, y), (t, d) in zip(pos, not_items):
    add_card(s, Inches(x), Inches(y), Inches(5.6), Inches(2.0), WHITE)
    add_rect(s, Inches(x), Inches(y), Inches(0.14), Inches(2.0), GOLD)
    _, ttf = add_textbox(s, Inches(x + 0.45), Inches(y + 0.3), Inches(4.9), Inches(0.5))
    set_run(ttf.paragraphs[0].add_run(), t, 18, True, RED_DARK)
    _, dtf = add_textbox(s, Inches(x + 0.45), Inches(y + 0.85), Inches(4.95), Inches(1.0))
    set_run(dtf.paragraphs[0].add_run(), d, 13.5, False, INK)


# ---------- 第 11 页：正确选择 ----------
s = prs.slides.add_slide(BLANK)
fill_bg(s, CREAM)
content_header(s, "正确选择 · 在奋斗与从容间找到平衡", "THE RIGHT PATH", 11)

points = [
    "奋斗是青春最亮丽的底色——人生因奋斗而精彩。",
    "拒绝无效内卷，追求高效成长——让努力更有意义。",
    "避免消极躺平，保持向上姿态——让心态始终在线。",
    "找到自己的节奏，实现有意义的人生——不盲从、不虚度。",
]
_, ptf = add_textbox(s, Inches(1.0), Inches(1.85), Inches(11.3), Inches(3.6))
for i, ptext in enumerate(points):
    # 序号
    add_card(s, Inches(0.95), Inches(1.85) + Inches(0.85) * i, Inches(0.5), Inches(0.5), RED)
    nb = s.shapes[-1]
    nb.text_frame.word_wrap = False
    nbp = nb.text_frame.paragraphs[0]
    nbp.alignment = PP_ALIGN.CENTER
    set_run(nbp.add_run(), str(i + 1), 16, True, WHITE)
    _, ttf = add_textbox(s, Inches(1.7), Inches(1.82) + Inches(0.85) * i, Inches(10.3), Inches(0.75))
    tp = ttf.paragraphs[0]
    set_run(tp.add_run(), ptext, 15.5, False, INK)

# 底部金句
add_card(s, Inches(0.9), Inches(5.5), Inches(11.5), Inches(1.25), WHITE)
add_rect(s, Inches(0.9), Inches(5.5), Inches(0.12), Inches(1.25), GOLD)
_, qtf = add_textbox(s, Inches(1.25), Inches(5.7), Inches(10.9), Inches(0.9))
qp = qtf.paragraphs[0]
set_run(qp.add_run(), "青春由磨砺而出彩，人生因奋斗而升华。", 19, True, RED_DARK)


# ---------- 第 12 页：结语 ----------
s = prs.slides.add_slide(BLANK)
fill_bg(s, RED)
add_rect(s, Inches(0), Inches(0), SLIDE_W, Inches(0.16), GOLD)
add_rect(s, Inches(0), Inches(7.34), SLIDE_W, Inches(0.16), GOLD)
add_rect(s, Inches(0), Inches(0.16), Inches(0.28), Inches(7.18), RED_DARK)

_, htf = add_textbox(s, Inches(1.0), Inches(1.1), Inches(11.3), Inches(0.6))
set_run(htf.paragraphs[0].add_run(), "结语 · 与青年共勉", 22, True, GOLD_LT)

_, btf = add_textbox(s, Inches(1.0), Inches(2.1), Inches(11.3), Inches(3.2))
lines = [
    "内卷不可怕，可怕的是迷失方向的盲目消耗；",
    "躺平不可取，可贵的是张弛有度的清醒选择。",
]
for i, ln in enumerate(lines):
    p = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
    p.space_after = Pt(12)
    set_run(p.add_run(), ln, 20, False, WHITE)

_, gtf = add_textbox(s, Inches(1.0), Inches(4.4), Inches(11.3), Inches(1.6))
gp = gtf.paragraphs[0]
gp.alignment = PP_ALIGN.CENTER
set_run(gp.add_run(), "愿每一位青年：", 18, True, GOLD_LT)
gp2 = gtf.add_paragraph()
gp2.alignment = PP_ALIGN.CENTER
gp2.space_before = Pt(10)
set_run(gp2.add_run(), "在可为之事上全力以赴，在不可为之事上保持清醒，", 16, False, WHITE)
gp3 = gtf.add_paragraph()
gp3.alignment = PP_ALIGN.CENTER
gp3.space_before = Pt(4)
set_run(gp3.add_run(), "拒绝无效内卷，不做消极躺平，书写无愧于时代的青春篇章。", 16, False, WHITE)

# 讨论环节
add_rect(s, Inches(3.17), Inches(6.2), Inches(7.0), Inches(0.02), GOLD)
_, dtf = add_textbox(s, Inches(1.0), Inches(6.4), Inches(11.3), Inches(0.6))
dp = dtf.paragraphs[0]
dp.alignment = PP_ALIGN.CENTER
set_run(dp.add_run(), "课堂讨论：你如何看待身边的「内卷」与「躺平」？", 14, False, GOLD_LT)


out = "躺平与内卷——青年学生的为与不为.pptx"
prs.save(out)
print("saved:", out, "| slides:", len(prs.slides.__iter__.__self__._sldIdLst))
